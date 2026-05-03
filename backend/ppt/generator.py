from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import io
import urllib.request

# 配色
C = {
    "blue":     RGBColor(0x25, 0x63, 0xEB),
    "blue_d":   RGBColor(0x1D, 0x4E, 0xD8),
    "blue_l":   RGBColor(0xDB, 0xEA, 0xFE),
    "green":    RGBColor(0x05, 0x96, 0x69),
    "orange":   RGBColor(0xEA, 0x58, 0x0C),
    "purple":   RGBColor(0x7C, 0x3A, 0xED),
    "dark":     RGBColor(0x11, 0x18, 0x27),
    "gray":     RGBColor(0x6B, 0x72, 0x80),
    "gray_l":   RGBColor(0xF3, 0xF4, 0xF6),
    "white":    RGBColor(0xFF, 0xFF, 0xFF),
}


def generate_from_dict(data: dict, output_path: str) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide_data in data.get("slides", []):
        if _is_empty_slide(slide_data):
            continue
        layout = slide_data.get("layout", slide_data.get("type", "content"))
        if layout in ("title", "cover"):
            _add_title_slide(prs, slide_data)
        elif layout == "data":
            _add_data_slide(prs, slide_data)
        elif layout == "summary":
            _add_summary_slide(prs, slide_data)
        else:
            _add_content_slide(prs, slide_data)

    prs.save(output_path)
    return output_path


def _is_empty_slide(data: dict) -> bool:
    title = data.get("title", "").strip()
    bullets = data.get("bullets", data.get("bullet", []))
    points = data.get("key_points", data.get("keypoints", []))
    desc = data.get("description", "").strip()
    subtitle = data.get("subtitle", "").strip()

    layout = data.get("layout", data.get("type", ""))
    if layout in ("title", "cover"):
        return not title

    if isinstance(bullets, str):
        bullets = [bullets]
    if isinstance(points, str):
        points = [points]
    has_bullets = bullets and any(str(b).strip() and not _is_placeholder(b) for b in bullets)
    has_points = points and any(str(p).strip() and not _is_placeholder(p) for p in points)
    return not has_bullets and not has_points and not desc


def _is_placeholder(text) -> bool:
    t = str(text).strip().lower()
    placeholders = ["描述", "要点", "内容", "数据", "结论", "placeholder", "todo", "tbd", "n/a"]
    if t in placeholders:
        return True
    if len(t) <= 4 and any(p in t for p in placeholders):
        return True
    return False


def _rect(slide, left, top, w, h, color):
    s = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def _circle(slide, left, top, size, color, alpha=0):
    s = slide.shapes.add_shape(9, Inches(left), Inches(top), Inches(size), Inches(size))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def _text_box(slide, left, top, w, h, text, size=18, color=None, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color or C["dark"]
    p.alignment = align
    return box


def _add_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _rect(slide, 0, 0, 8.5, 7.5, C["blue"])
    _rect(slide, 8.5, 0, 4.833, 7.5, C["gray_l"])

    _circle(slide, 7.0, -0.5, 2.5, C["blue_d"])
    _circle(slide, 10.5, 5.5, 3.0, C["blue_l"])

    title = data.get("title", "")
    _text_box(slide, 0.8, 2.0, 7.0, 2.0, title, size=44, color=C["white"], bold=True)

    subtitle = data.get("subtitle", "")
    if subtitle:
        _text_box(slide, 9.0, 3.0, 3.8, 2.0, subtitle, size=18, color=C["gray"])

    _rect(slide, 0.8, 4.3, 2.0, 0.06, C["orange"])


def _add_content_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _rect(slide, 0, 0, 13.333, 0.12, C["blue"])

    _rect(slide, 0.5, 0.4, 12.333, 0.9, C["gray_l"])
    _rect(slide, 0.5, 0.4, 0.08, 0.9, C["blue"])
    _text_box(slide, 0.9, 0.55, 11, 0.7, data.get("title", ""), size=28, bold=True)

    bullets = data.get("bullets", data.get("bullet", []))
    if not bullets:
        return
    if isinstance(bullets, str):
        bullets = [bullets]

    y = 1.7
    for i, bullet in enumerate(bullets):
        if not str(bullet).strip() or _is_placeholder(bullet):
            continue

        _rect(slide, 0.7, y, 11.9, 0.95, C["white"] if i % 2 == 0 else C["gray_l"])

        num_circle = _circle(slide, 0.85, y + 0.15, 0.6, C["blue"])
        _text_box(slide, 0.85, y + 0.18, 0.6, 0.55,
                  str(i + 1), size=16, color=C["white"], bold=True, align=PP_ALIGN.CENTER)

        _text_box(slide, 1.7, y + 0.15, 10.5, 0.7, str(bullet), size=17)

        y += 1.1
        if y > 6.2:
            break


def _add_data_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _rect(slide, 0, 0, 13.333, 0.12, C["green"])
    _rect(slide, 0.5, 0.4, 12.333, 0.9, C["gray_l"])
    _rect(slide, 0.5, 0.4, 0.08, 0.9, C["green"])
    _text_box(slide, 0.9, 0.55, 11, 0.7, data.get("title", ""), size=28, bold=True)

    desc = data.get("description", "")
    if desc:
        _text_box(slide, 0.9, 1.5, 11.5, 0.6, desc, size=16, color=C["gray"])

    key_points = data.get("key_points", data.get("keypoints", []))
    if not key_points:
        return
    if isinstance(key_points, str):
        key_points = [key_points]

    cols = min(len(key_points), 3)
    card_w = (11.5 - 0.3 * (cols - 1)) / cols
    y_start = 2.4

    for i, point in enumerate(key_points):
        if not str(point).strip() or _is_placeholder(point):
            continue
        col = i % cols
        row = i // cols
        x = 0.7 + col * (card_w + 0.3)
        y = y_start + row * 1.8

        if y > 6.0:
            break

        _rect(slide, x, y, card_w, 1.5, C["gray_l"])
        _rect(slide, x, y, card_w, 0.06, C["green"])

        _text_box(slide, x + 0.2, y + 0.2, card_w - 0.4, 1.2, str(point), size=16)


def _add_summary_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    _rect(slide, 0, 0, 13.333, 7.5, C["blue"])

    _circle(slide, -1, -1, 4, C["blue_d"])
    _circle(slide, 11, 5.5, 3.5, C["blue_d"])

    _text_box(slide, 1, 0.6, 11.333, 1.0,
              data.get("title", "总结"), size=36, color=C["white"], bold=True, align=PP_ALIGN.CENTER)

    _rect(slide, 5.5, 1.7, 2.333, 0.04, C["orange"])

    bullets = data.get("bullets", data.get("bullet", []))
    if isinstance(bullets, str):
        bullets = [bullets]
    y = 2.2
    for bullet in bullets:
        if not str(bullet).strip() or _is_placeholder(bullet):
            continue
        _text_box(slide, 1.5, y, 10.333, 0.8,
                  f"✓  {bullet}", size=20, color=C["white"])
        y += 0.9
        if y > 6.5:
            break
